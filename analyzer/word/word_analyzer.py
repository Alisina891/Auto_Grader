import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


# ============================================================
# NORMALIZATION
# ============================================================

def normalize_text(value):
    """
    Normalize text for flexible comparison.

    Examples:
        Introduction
        introduction
        INTRODUCTION
        "  Introduction  "

    are treated as the same normalized text.
    """

    if value is None:
        return ""

    text = str(value)

    text = (
        text
        .replace("\u200b", "")
        .replace("\u200c", "")
        .replace("\u200d", "")
        .replace("\u200f", "")
        .replace("\ufeff", "")
        .replace("\u00a0", " ")
    )

    text = " ".join(text.strip().split())

    return text.casefold()


def normalize_heading_text(value):
    return normalize_text(value)


# ============================================================
# DOCUMENT INFO
# ============================================================

def get_document_info(filepath):
    document = Document(filepath)

    return {
        "file_name": os.path.basename(filepath),
        "document_type": "docx",
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "section_count": len(document.sections),
    }


# ============================================================
# ALIGNMENT
# ============================================================

def get_alignment_name(alignment):

    if alignment is None:
        return None

    alignment_map = {
        WD_ALIGN_PARAGRAPH.LEFT: "left",
        WD_ALIGN_PARAGRAPH.CENTER: "center",
        WD_ALIGN_PARAGRAPH.RIGHT: "right",
        WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
        WD_ALIGN_PARAGRAPH.DISTRIBUTE: "distribute",
    }

    return alignment_map.get(
        alignment,
        str(alignment)
    )


# ============================================================
# RUN FORMATTING
# ============================================================

def extract_run_format(run):
    """
    Extract run formatting.

    IMPORTANT:
    Font size is intentionally NOT extracted.

    Font size will NEVER affect the score.
    """

    font = run.font

    return {
        "font_name": font.name,

        "bold": (
            bool(font.bold)
            if font.bold is not None
            else False
        ),

        "italic": (
            bool(font.italic)
            if font.italic is not None
            else False
        ),

        "underline": (
            bool(font.underline)
            if font.underline is not None
            else False
        ),
    }


# ============================================================
# PARAGRAPH FORMATTING
# ============================================================

def extract_paragraph_formatting(paragraph):

    formatting = {
        "alignment": get_alignment_name(
            paragraph.alignment
        ),

        "style": None,

        "left_indent": None,
        "right_indent": None,
        "first_line_indent": None,

        "space_before": None,
        "space_after": None,

        "line_spacing": None,
    }

    try:
        formatting["style"] = paragraph.style.name
    except Exception:
        pass

    try:

        fmt = paragraph.paragraph_format

        if fmt.left_indent is not None:
            formatting["left_indent"] = (
                fmt.left_indent.pt
            )

        if fmt.right_indent is not None:
            formatting["right_indent"] = (
                fmt.right_indent.pt
            )

        if fmt.first_line_indent is not None:
            formatting["first_line_indent"] = (
                fmt.first_line_indent.pt
            )

        if fmt.space_before is not None:
            formatting["space_before"] = (
                fmt.space_before.pt
            )

        if fmt.space_after is not None:
            formatting["space_after"] = (
                fmt.space_after.pt
            )

        if fmt.line_spacing is not None:
            formatting["line_spacing"] = (
                fmt.line_spacing
            )

    except Exception:
        pass

    return formatting


# ============================================================
# PARAGRAPH
# ============================================================

def extract_paragraph(paragraph, index):

    text = paragraph.text or ""

    style_name = None

    try:
        style_name = paragraph.style.name
    except Exception:
        pass

    runs = []

    for run in paragraph.runs:

        run_text = run.text or ""

        if not run_text:
            continue

        runs.append({
            "text": run_text,

            "normalized_text":
                normalize_text(run_text),

            "format":
                extract_run_format(run),
        })

    # --------------------------------------------------------
    # Detect Heading
    # --------------------------------------------------------

    is_heading = False
    heading_level = None

    if style_name:

        match = re.match(
            r"Heading\s+(\d+)",
            style_name,
            re.IGNORECASE
        )

        if match:

            is_heading = True

            heading_level = int(
                match.group(1)
            )

    # --------------------------------------------------------
    # Also detect title-like paragraphs
    # --------------------------------------------------------

    if (
        not is_heading
        and text.strip()
        and style_name
        and normalize_text(style_name) == "title"
    ):

        is_heading = True
        heading_level = 1

    return {

        "index": index,

        "text": text,

        "normalized_text":
            normalize_text(text),

        "style": style_name,

        "is_heading": is_heading,

        "heading_level": heading_level,

        "alignment":
            get_alignment_name(
                paragraph.alignment
            ),

        "runs": runs,

        "formatting":
            extract_paragraph_formatting(
                paragraph
            ),
    }


# ============================================================
# HEADINGS
# ============================================================

def extract_headings(paragraphs):

    headings = []

    for paragraph in paragraphs:

        if not paragraph["is_heading"]:
            continue

        text = paragraph["text"].strip()

        if not text:
            continue

        headings.append({

            "text": text,

            "normalized_text":
                paragraph["normalized_text"],

            "level":
                paragraph["heading_level"],

            "style":
                paragraph["style"],

            "paragraph_index":
                paragraph["index"],
        })

    return headings


# ============================================================
# SECTIONS
# ============================================================

def extract_sections(document):

    sections = []

    for index, section in enumerate(
        document.sections
    ):

        sections.append({

            "index": index,

            "start_type":
                str(section.start_type)
                if section.start_type is not None
                else None,

            "page_width":
                section.page_width.pt
                if section.page_width is not None
                else None,

            "page_height":
                section.page_height.pt
                if section.page_height is not None
                else None,

            "left_margin":
                section.left_margin.pt
                if section.left_margin is not None
                else None,

            "right_margin":
                section.right_margin.pt
                if section.right_margin is not None
                else None,

            "top_margin":
                section.top_margin.pt
                if section.top_margin is not None
                else None,

            "bottom_margin":
                section.bottom_margin.pt
                if section.bottom_margin is not None
                else None,
        })

    return sections


# ============================================================
# TABLE CELL
# ============================================================

def extract_table_cell(cell):

    text = cell.text or ""

    return {

        "text": text,

        "normalized_text":
            normalize_text(text),
    }


# ============================================================
# TABLES
# ============================================================

def extract_tables(document):

    tables = []

    for table_index, table in enumerate(
        document.tables
    ):

        rows = []

        for row_index, row in enumerate(
            table.rows
        ):

            cells = []

            for column_index, cell in enumerate(
                row.cells
            ):

                cells.append({

                    "index": column_index,

                    **extract_table_cell(cell)
                })

            rows.append({

                "index": row_index,

                "cells": cells,
            })

        column_count = 0

        if rows:

            column_count = max(
                len(row["cells"])
                for row in rows
            )

        headers = []

        if rows:

            headers = [
                cell["text"]
                for cell in rows[0]["cells"]
            ]

        tables.append({

            "index": table_index,

            "row_count":
                len(rows),

            "column_count":
                column_count,

            "headers":
                headers,

            "normalized_headers": [
                normalize_text(header)
                for header in headers
            ],

            "rows":
                rows,
        })

    return tables


# ============================================================
# IMAGES
# ============================================================

def extract_images(document):

    images = []

    relationships = document.part.rels

    for relationship_id, relationship in (
        relationships.items()
    ):

        if "image" not in relationship.reltype:
            continue

        images.append({

            "relationship_id":
                relationship_id,

            "target":
                relationship.target_ref,

            # Filename is informational only.
            # It MUST NOT affect grading.
            "filename":
                os.path.basename(
                    relationship.target_ref
                ),
        })

    return images


# ============================================================
# SHAPES
# ============================================================

def extract_shapes(document):
    """
    Detect Word shapes / drawing objects.

    This does not attempt to judge whether a shape is
    visually identical.

    It only identifies the presence and general type
    of drawing objects.

    Shape filename/location is not used for grading.
    """

    shapes = []

    try:

        root = document.part.element

        namespaces = {
            "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
            "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
            "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
            "wps": "http://schemas.microsoft.com/office/word/2010/wordprocessingShape",
            "v": "urn:schemas-microsoft-com:vml",
        }

        # ----------------------------------------------------
        # Drawing objects
        # ----------------------------------------------------

        drawings = root.xpath(
            ".//w:drawing"
        )

        for drawing in drawings:

            shape_type = "drawing"

            if drawing.xpath(
                ".//wps:wsp",
                namespaces=namespaces
            ):
                shape_type = "shape"

            elif drawing.xpath(
                ".//a:graphic",
                namespaces=namespaces
            ):
                shape_type = "graphic"

            elif drawing.xpath(
                ".//wp:inline",
                namespaces=namespaces
            ):
                shape_type = "inline_object"

            elif drawing.xpath(
                ".//wp:anchor",
                namespaces=namespaces
            ):
                shape_type = "anchored_object"

            shapes.append({

                "index":
                    len(shapes),

                "type":
                    shape_type,
            })

        # ----------------------------------------------------
        # Legacy VML shapes
        # ----------------------------------------------------

        legacy_shapes = root.xpath(
            ".//w:pict"
        )

        for shape in legacy_shapes:

            shapes.append({

                "index":
                    len(shapes),

                "type":
                    "legacy_shape",
            })

    except Exception:

        pass

    return shapes


# ============================================================
# TEXT CONTENT
# ============================================================

def extract_text_content(paragraphs):

    content = []

    for paragraph in paragraphs:

        text = paragraph["text"].strip()

        if not text:
            continue

        content.append({

            "text": text,

            "normalized_text":
                paragraph["normalized_text"],

            "paragraph_index":
                paragraph["index"],

            "is_heading":
                paragraph["is_heading"],

            "heading_level":
                paragraph["heading_level"],
        })

    return content


# ============================================================
# FORMATTING SUMMARY
# ============================================================

def extract_formatting_summary(paragraphs):

    font_names = []

    bold_runs = 0
    italic_runs = 0
    underline_runs = 0

    alignments = []

    for paragraph in paragraphs:

        alignment = paragraph.get(
            "alignment"
        )

        if alignment:
            alignments.append(
                alignment
            )

        for run in paragraph["runs"]:

            font_name = run["format"].get(
                "font_name"
            )

            if font_name:
                font_names.append(
                    font_name
                )

            if run["format"].get("bold"):
                bold_runs += 1

            if run["format"].get("italic"):
                italic_runs += 1

            if run["format"].get("underline"):
                underline_runs += 1

    unique_fonts = []

    for font in font_names:

        if font not in unique_fonts:

            unique_fonts.append(font)

    unique_alignments = []

    for alignment in alignments:

        if alignment not in unique_alignments:

            unique_alignments.append(
                alignment
            )

    return {

        "font_names":
            unique_fonts,

        "bold_run_count":
            bold_runs,

        "italic_run_count":
            italic_runs,

        "underline_run_count":
            underline_runs,

        "alignments":
            unique_alignments,

        # NEVER grade font size.
        "font_size_checked":
            False,
    }


# ============================================================
# CONTENT CANDIDATES
# ============================================================

def extract_content_candidates(paragraphs):

    candidates = []

    for paragraph in paragraphs:

        text = paragraph["text"].strip()

        if not text:
            continue

        candidates.append({

            "text":
                text,

            "normalized_text":
                paragraph["normalized_text"],

            "paragraph_index":
                paragraph["index"],

            "is_heading":
                paragraph["is_heading"],

            "heading_level":
                paragraph["heading_level"],
        })

    return candidates


# ============================================================
# FULL RULES
# ============================================================

def extract_full_rules(filepath):
    """
    Extract information from a Word document.

    IMPORTANT DESIGN:

    The analyzer only extracts facts.

    It does NOT decide:
        - what is correct
        - what is wrong
        - what is required
        - what score should be given

    The Rule Engine makes those decisions.

    FLEXIBILITY:

    - Heading text is informational.
    - Paragraph text is informational.
    - Paragraph position is informational.
    - Table position is informational.
    - Image filename is informational.
    - Font size is completely ignored.
    """

    document = Document(filepath)

    # --------------------------------------------------------
    # Paragraphs
    # --------------------------------------------------------

    paragraphs = []

    for index, paragraph in enumerate(
        document.paragraphs
    ):

        paragraphs.append(
            extract_paragraph(
                paragraph,
                index
            )
        )

    # --------------------------------------------------------
    # Headings
    # --------------------------------------------------------

    headings = extract_headings(
        paragraphs
    )

    # --------------------------------------------------------
    # Sections
    # --------------------------------------------------------

    sections = extract_sections(
        document
    )

    # --------------------------------------------------------
    # Tables
    # --------------------------------------------------------

    tables = extract_tables(
        document
    )

    # --------------------------------------------------------
    # Images
    # --------------------------------------------------------

    images = extract_images(
        document
    )

    # --------------------------------------------------------
    # Shapes
    # --------------------------------------------------------

    shapes = extract_shapes(
        document
    )

    # --------------------------------------------------------
    # Text
    # --------------------------------------------------------

    text_content = extract_text_content(
        paragraphs
    )

    # --------------------------------------------------------
    # Candidates
    # --------------------------------------------------------

    content_candidates = (
        extract_content_candidates(
            paragraphs
        )
    )

    # --------------------------------------------------------
    # Formatting
    # --------------------------------------------------------

    formatting = extract_formatting_summary(
        paragraphs
    )

    # --------------------------------------------------------
    # Rules
    # --------------------------------------------------------

    return {

        "file_name":
            os.path.basename(filepath),

        "document_type":
            "docx",

        "paragraph_count":
            len(paragraphs),

        "heading_count":
            len(headings),

        "table_count":
            len(tables),

        "image_count":
            len(images),

        "shape_count":
            len(shapes),

        "section_count":
            len(sections),

        "paragraphs":
            paragraphs,

        "headings":
            headings,

        "sections":
            sections,

        "tables":
            tables,

        "images":
            images,

        "shapes":
            shapes,

        "text_content":
            text_content,

        "content_candidates":
            content_candidates,

        "formatting":
            formatting,

        "comparison_settings": {

            # Flexible by default.
            "paragraph_location_strict":
                False,

            "heading_location_strict":
                False,

            "table_location_strict":
                False,

            # Never enable this.
            "font_size_strict":
                False,

            # Image filename never matters.
            "image_filename_strict":
                False,

            # Shape exact type is not required
            # unless explicitly configured later.
            "shape_type_strict":
                False,
        },
    }


# ============================================================
# PRINT SUMMARY
# ============================================================

def print_rule_summary(rules):

    print()
    print("=" * 70)
    print("WORD DOCUMENT ANALYSIS")
    print("=" * 70)

    print()
    print(
        f"File: {rules['file_name']}"
    )

    print(
        f"Paragraphs: "
        f"{rules['paragraph_count']}"
    )

    print(
        f"Headings: "
        f"{rules['heading_count']}"
    )

    print(
        f"Tables: "
        f"{rules['table_count']}"
    )

    print(
        f"Images: "
        f"{rules['image_count']}"
    )

    print(
        f"Shapes: "
        f"{rules['shape_count']}"
    )

    print(
        f"Sections: "
        f"{rules['section_count']}"
    )

    print()
    print("-" * 70)
    print("HEADINGS")
    print("-" * 70)

    for heading in rules["headings"]:

        print(
            f"  H{heading['level']}: "
            f"{heading['text']}"
        )

    print()
    print("-" * 70)
    print("TABLES")
    print("-" * 70)

    for table in rules["tables"]:

        print(
            f"  Table {table['index'] + 1}: "
            f"{table['row_count']} rows x "
            f"{table['column_count']} columns"
        )

        print(
            f"    Headers: "
            f"{table['headers']}"
        )

    print()
    print("-" * 70)
    print("IMAGES")
    print("-" * 70)

    for image in rules["images"]:

        print(
            f"  - {image['filename']}"
        )

    print()
    print("-" * 70)
    print("SHAPES")
    print("-" * 70)

    for shape in rules["shapes"]:

        print(
            f"  - {shape['type']}"
        )

    print()
    print("-" * 70)
    print("FORMATTING")
    print("-" * 70)

    formatting = rules["formatting"]

    print(
        f"Font names: "
        f"{formatting['font_names']}"
    )

    print(
        f"Bold runs: "
        f"{formatting['bold_run_count']}"
    )

    print(
        f"Italic runs: "
        f"{formatting['italic_run_count']}"
    )

    print(
        f"Underline runs: "
        f"{formatting['underline_run_count']}"
    )

    print(
        f"Alignments: "
        f"{formatting['alignments']}"
    )

    print(
        "Font size checking: DISABLED"
    )

    print()
    print("=" * 70)
    print("ANALYSIS COMPLETE")
    print("=" * 70)


# ============================================================
# TEST
# ============================================================

if __name__ == "__main__":

    base_directory = os.path.dirname(
        os.path.abspath(__file__)
    )

    filepath = os.path.join(
        base_directory,
        "master_project.docx"
    )

    if not os.path.exists(filepath):

        print()
        print(
            f"File not found: {filepath}"
        )

        print()
        print(
            "Put master_project.docx inside:"
        )

        print(
            base_directory
        )

    else:

        rules = extract_full_rules(
            filepath
        )

        print_rule_summary(
            rules
        )